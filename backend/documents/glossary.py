"""Optional Excel-backed terminology and calculation templates for query expansion."""
from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path

from . import config


_STOP = {"a", "an", "and", "as", "at", "by", "for", "from", "in", "is", "it", "of", "on",
         "or", "the", "this", "to", "used", "with", "metric", "metrics", "page", "period",
         "actual", "compare", "earlier", "estimate", "figure", "figures", "full", "given",
         "latest", "later", "month", "months", "planned", "shows", "total", "value", "values",
         "year", "years", "fy"}


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _tokens(value: str) -> set[str]:
    words = [x for x in re.findall(r"[a-z0-9%]+", value.lower())
             if len(x) > 1 and x not in _STOP and not x.isdigit()]
    return {x[:5] for x in words}


@lru_cache(maxsize=4)
def _read(path_text: str, sheet_name: str, modified_ns: int) -> tuple[dict, ...]:
    # modified_ns is part of the cache key so an edited glossary is picked up without restart.
    del modified_ns
    from openpyxl import load_workbook

    wb = load_workbook(path_text, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        return ()
    ws = wb[sheet_name]
    entries: list[dict] = []
    mode = ""
    for row in ws.iter_rows(values_only=True):
        values = [_norm(v) for v in row]
        if values[:2] == ["Category", "Term / Metric"]:
            mode = "term"
            continue
        if values[:2] == ["Page (Category)", "Derived Metric"]:
            mode = "formula"
            continue
        if mode == "term" and len(values) >= 6 and values[1]:
            entries.append({"kind": "term", "category": values[0], "name": values[1],
                            "alias": values[2], "definition": values[3], "why": values[4],
                            "formula": values[5]})
        elif mode == "formula" and len(values) >= 6 and values[1]:
            # Keep the category text but remove its page number: the glossary guides terminology,
            # while retrieval must still discover the source page from the user's question.
            category = re.sub(r"(?i)^page\s+\d+\s*[-–—:]?\s*", "", values[0])
            entries.append({"kind": "formula", "category": category, "name": values[1],
                            "alias": values[2], "definition": values[3], "why": values[4],
                            "formula": values[5]})
    return tuple(entries)


def entries() -> tuple[dict, ...]:
    path = config.GLOSSARY_FILE
    if not path:
        return ()
    path = Path(path)
    if not path.is_file():
        return ()
    return _read(str(path.resolve()), config.GLOSSARY_SHEET, path.stat().st_mtime_ns)


def _term_hit(entry: dict, question: str, qtokens: set[str]) -> bool:
    names = [entry["name"], entry["alias"]]
    for name in names:
        for variant in re.split(r"\s*/\s*|\s+or\s+", name, flags=re.I):
            phrase = " ".join(re.findall(r"[a-z0-9%]+", variant.lower()))
            if phrase and re.search(rf"(?<![a-z0-9]){re.escape(phrase)}(?![a-z0-9])", question.lower()):
                return True
    # Full-form matches need at least two distinctive words; one broad word such as "cost" is
    # not enough to inject unrelated glossary content.
    return len(_tokens(entry["alias"]) & qtokens) >= 2


def _formula_score(entry: dict, qtokens: set[str]) -> int:
    key = _tokens(entry["name"] + " " + entry["category"] + " " + entry["definition"])
    overlap = key & qtokens
    return len(overlap) if len(overlap) >= 2 else 0


def match(question: str, limit: int = 5) -> list[dict]:
    qtokens = _tokens(question)
    terms = [e for e in entries() if e["kind"] == "term" and _term_hit(e, question, qtokens)]
    formulas = sorted((( _formula_score(e, qtokens), e) for e in entries() if e["kind"] == "formula"),
                      key=lambda pair: pair[0], reverse=True)
    picked = terms + [e for score, e in formulas if score > 0]
    # Stable deduplication and a short expansion keep small-model prompts focused.
    return list({(e["kind"], e["name"]): e for e in picked}.values())[:limit]


def expand(question: str, matches: list[dict]) -> str:
    if not matches:
        return question
    additions = []
    for e in matches:
        additions.extend(x for x in (e["name"], e["alias"], e["category"]) if x)
    return question + "\nRelevant terminology: " + "; ".join(dict.fromkeys(additions))


def context(matches: list[dict]) -> str:
    if not matches:
        return ""
    lines = ["\nGlossary templates (terminology/formulas only; not report evidence or values):"]
    for e in matches:
        detail = f"- {e['name']}"
        if e["alias"]:
            detail += f" ({e['alias']})"
        if e["definition"]:
            detail += f": {e['definition']}"
        if e["formula"] and "reported directly" not in e["formula"].lower():
            detail += f" Formula template: {e['formula']}"
        lines.append(detail)
    return "\n".join(lines)
